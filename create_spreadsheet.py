import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_aliexpress_pricing_table():
    wb = openpyxl.Workbook()
    
    # 1. Calculator Sheet
    ws = wb.active
    ws.title = "定价计算器"
    
    # Define styles
    header_font = Font(name='微软雅黑', bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    input_fill = PatternFill(start_color='EBF1DE', end_color='EBF1DE', fill_type='solid') # Light green for inputs
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align_center = Alignment(horizontal='center', vertical='center')
    
    # Inputs section
    ws['A1'] = "基础参数设置"
    ws['A1'].font = Font(name='微软雅黑', size=14, bold=True)
    
    inputs = [
        ("产品成本 (RMB)", 13.5),
        ("产品重量 (g)", 100),
        ("美元汇率 (USD/RMB)", 7.25),
        ("目标利润率 (%)", 0.2),
        ("活动折扣 (%)", 0.3),
        ("平台佣金 (%)", 0.08),
        ("联盟营销/其他 (%)", 0.05)
    ]
    
    for i, (label, val) in enumerate(inputs, start=2):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=val)
        cell.fill = input_fill
        cell.border = border
        if "%" in label:
            cell.number_format = '0%'
        elif "汇率" in label:
            cell.number_format = '0.00'
        else:
            cell.number_format = '0.00'

    # Results Table Headers
    table_start_row = 11
    headers = ["国家", "物流方式", "运费 (RMB)", "总成本 (RMB)", "建议售价 (USD)", "折后标价 (USD)", "预估利润 (RMB)", "利润率"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=table_start_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = align_center

    # 2. Shipping Rates Sheet
    ws_rates = wb.create_sheet("运费数据库")
    ws_rates.append(["国家", "物流名称", "计费重(KG)单价", "挂号费/处理费"])
    
    # Latest extracted data (Samples)
    rates_data = [
        ["俄罗斯 (RU)", "无忧标准", 136, 21], # Approximate from general trend
        ["美国 (US)", "无忧标准", 145, 25],
        ["法国 (FR)", "无忧标准", 91.74, 17.02],
        ["德国 (DE)", "无忧标准", 70.44, 18.85],
        ["英国 (GB)", "无忧标准", 67.76, 15.9],
        ["西班牙 (ES)", "无忧标准", 64.67, 19.26],
        ["澳大利亚 (AU)", "无忧标准", 89.52, 18.4],
        ["加拿大 (CA)", "无忧标准", 122.44, 27.6],
        ["日本 (JP)", "无忧标准", 64.45, 21.2],
        ["韩国 (KR)", "无忧标准", 58.07, 10.09], # Estimate based on trend
        ["以色列 (IL)", "无忧标准", 128.01, 20.45],
        ["荷兰 (NL)", "无忧标准", 75.32, 18.45],
        ["意大利 (IT)", "无忧标准", 63.8, 23.4],
        ["波兰 (PL)", "无忧标准", 72.15, 13.12],
        ["乌克兰 (UA)", "无忧标准", 138.2, 18.7],
        ["巴西 (BR)", "无忧标准", 115.5, 35.5]
    ]
    
    for row in rates_data:
        ws_rates.append(row)
    
    # Add formulas to Calculator sheet
    num_countries = len(rates_data)
    for i in range(num_countries):
        curr_row = table_start_row + 1 + i
        rate_row = i + 2
        
        # 国家
        ws.cell(row=curr_row, column=1, value=f"='运费数据库'!A{rate_row}")
        # 物流方式
        ws.cell(row=curr_row, column=2, value=f"='运费数据库'!B{rate_row}")
        
        # 运费 (RMB) = (Weight/1000) * Rate_KG + Handling
        weight_ref = "$B$3"
        ws.cell(row=curr_row, column=3, value=f"=({weight_ref}/1000)*'运费数据库'!C{rate_row}+'运费数据库'!D{rate_row}")
        
        # 总成本 (RMB) = Product Cost + Shipping
        cost_ref = "$B$2"
        ws.cell(row=curr_row, column=4, value=f"={cost_ref}+C{curr_row}")
        
        # 建议售价 (USD) = TotalCost / ExchangeRate / (1 - Margin - Commission - Other)
        # Note: B4=Exchange, B5=Margin, B7=Commission, B8=Other
        ws.cell(row=curr_row, column=5, value=f"=(D{curr_row}/$B$4)/(1-$B$5-$B$7-$B$8)")
        
        # 折后标价 (USD) = SellingPrice / (1 - Discount)
        # B6=Discount
        ws.cell(row=curr_row, column=6, value=f"=E{curr_row}/(1-$B$6)")
        
        # 预估利润 (RMB) = SellingPrice * ExchangeRate * (1 - Commission - Other) - TotalCost
        ws.cell(row=curr_row, column=7, value=f"=E{curr_row}*$B$4*(1-$B$7-$B$8)-D{curr_row}")
        
        # 利润率 = Profit / (SellingPrice * ExchangeRate)
        ws.cell(row=curr_row, column=8, value=f"=G{curr_row}/(E{curr_row}*$B$4)")

    # Formatting results
    for r in range(table_start_row + 1, table_start_row + 1 + num_countries):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if c in [3, 4, 7]:
                cell.number_format = '¥#,##0.00'
            elif c in [5, 6]:
                cell.number_format = '$#,##0.00'
            elif c == 8:
                cell.number_format = '0.0%'

    # Adjust column widths
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 15
    ws.column_dimensions['A'].width = 20

    # Save
    save_path = r"C:\Users\Administrator\Desktop\速卖通运费利润计算表_2026.xlsx"
    wb.save(save_path)
    return save_path

try:
    path = create_aliexpress_pricing_table()
    print(f"Success: {path}")
except Exception as e:
    print(f"Error: {e}")
